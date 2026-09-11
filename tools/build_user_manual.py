from pathlib import Path

from PIL import Image, ImageChops
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "docs" / "assets"
OUT = ROOT / "docs" / "01-user-manual.docx"
CROPS = ASSETS / "manual-crops"

BLUE = "0F766E"
BLUE_DARK = "134E4A"
SLATE = "334155"
MUTED = "64748B"
LIGHT_BLUE = "E6FFFA"
LIGHT_GRAY = "F1F5F9"
LIGHT_GOLD = "FFFBEB"
LIGHT_RED = "FEF2F2"
WHITE = "FFFFFF"
TABLE_WIDTH = 9360


def set_run_font(run, name="Microsoft YaHei", size=10.5, color=SLATE, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.first_child_found_in("w:shd")
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="CBD5E1", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), size)
        edge.set(qn("w:color"), color)


def set_table_geometry(table, widths):
    if sum(widths) != TABLE_WIDTH:
        raise ValueError(f"table widths must total {TABLE_WIDTH}: {widths}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.insert(0, tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def clear_cell(cell):
    cell.text = ""
    return cell.paragraphs[0]


def add_cell_text(cell, text, bold=False, color=SLATE, size=9.5):
    paragraph = clear_cell(cell)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, color=color, bold=bold)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    header = table.rows[0]
    repeat_header(header)
    for i, value in enumerate(headers):
        add_cell_text(header.cells[i], value, bold=True, color=BLUE_DARK, size=9.5)
        set_cell_shading(header.cells[i], "E8EEF5")
        set_cell_border(header.cells[i])
    for row_values in rows:
        row = table.add_row()
        for i, value in enumerate(row_values):
            add_cell_text(row.cells[i], value, size=9.2)
            set_cell_border(row.cells[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_step(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_note(doc, title, text, fill=LIGHT_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_WIDTH])
    repeat_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color=accent, size="10")
    paragraph = clear_cell(cell)
    paragraph.paragraph_format.space_after = Pt(2)
    title_run = paragraph.add_run(title)
    set_run_font(title_run, size=10, color=accent, bold=True)
    body_run = paragraph.add_run("  " + text)
    set_run_font(body_run, size=10, color=SLATE)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def crop_image(source, target):
    image = Image.open(source).convert("RGB")
    background = Image.new("RGB", image.size, image.getpixel((image.width - 1, image.height - 1)))
    diff = ImageChops.difference(image, background).convert("L")
    mask = diff.point(lambda value: 255 if value > 10 else 0)
    bbox = mask.getbbox()
    if bbox:
        left, top, right, bottom = bbox
        left = max(0, left - 8)
        top = max(0, top - 8)
        right = min(image.width, right + 8)
        bottom = min(image.height, bottom + 8)
        image = image.crop((left, top, right, bottom))
    image.save(target, optimize=True)


def add_screenshot(doc, filename, caption, alt):
    source = ASSETS / filename
    target = CROPS / filename
    target.parent.mkdir(parents=True, exist_ok=True)
    crop_image(source, target)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(3)
    shape = paragraph.add_run().add_picture(str(target), width=Inches(6.22))
    shape._inline.docPr.set("descr", alt)
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(10)
    run = caption_paragraph.add_run(caption)
    set_run_font(run, size=9, color=MUTED, italic=True)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(SLATE)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 11.5, BLUE_DARK, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def configure_page(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.38)
    section.footer_distance = Inches(0.38)
    header = section.header
    header_table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    header_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    header_table.autofit = False
    repeat_header(header_table.rows[0])
    header_table.columns[0].width = Inches(3.25)
    header_table.columns[1].width = Inches(3.25)
    for cell in header_table.rows[0].cells:
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
    left = clear_cell(header_table.cell(0, 0))
    left_run = left.add_run("综合项目管理系统")
    set_run_font(left_run, size=8.5, color=MUTED, bold=True)
    right = clear_cell(header_table.cell(0, 1))
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right_run = right.add_run("用户使用手册 | v3.0.4")
    set_run_font(right_run, size=8.5, color=MUTED)
    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer_paragraph.add_run("内部使用  |  ")
    set_run_font(footer_run, size=8.5, color=MUTED)
    add_page_number(footer_paragraph)


def build():
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    doc.core_properties.title = "综合项目管理系统用户使用手册"
    doc.core_properties.subject = "工时、项目、合同、收付款和看板操作说明"
    doc.core_properties.author = "综合项目管理系统"
    doc.core_properties.keywords = "工时管理,项目管理,合同,收付款,飞书"

    cover_kicker = doc.add_paragraph()
    cover_kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_kicker.paragraph_format.space_before = Pt(40)
    cover_kicker.paragraph_format.space_after = Pt(10)
    run = cover_kicker.add_run("USER GUIDE  |  v3.0.4")
    set_run_font(run, size=10, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(7)
    run = title.add_run("综合项目管理系统")
    set_run_font(run, size=28, color=BLUE_DARK, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(22)
    run = subtitle.add_run("用户使用手册")
    set_run_font(run, size=17, color=SLATE, bold=True)

    meta = doc.add_table(rows=3, cols=2)
    set_table_geometry(meta, [1800, 7560])
    repeat_header(meta.rows[0])
    meta_rows = [("适用版本", "v3.0.4"), ("更新时间", "2026 年 8 月 31 日"), ("适用人员", "员工、审批人、项目负责人、出纳和系统管理员")]
    for row, (label, value) in zip(meta.rows, meta_rows):
        add_cell_text(row.cells[0], label, bold=True, color=BLUE_DARK, size=9.5)
        add_cell_text(row.cells[1], value, size=9.5)
        set_cell_shading(row.cells[0], "E8EEF5")
        for cell in row.cells:
            set_cell_border(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_note(doc, "阅读提示", "本手册中的截图为 v3.0 原型界面示例，操作规则和菜单权限以当前部署系统为准。", fill=LIGHT_GOLD, accent="A16207")
    add_screenshot(doc, "manual-workbench.png", "图 1  工作台界面示例：项目负责人区与员工工时区", "工作台界面示例")
    doc.add_page_break()

    add_heading(doc, "目录", 1)
    for item in ("快速开始", "角色与菜单", "工作台", "工时填报", "工时审批", "项目管理", "合同与收付款", "工时看板与项目看板", "系统管理", "常见问题与安全提示"):
        add_step(doc, item)
    add_note(doc, "系统边界", "飞书用于免登、组织架构和应用消息通知；业务数据、校验、审批操作和经营流水保存在本系统。系统不直接发起银行转账，也不替代财务总账。", fill=LIGHT_BLUE)

    add_heading(doc, "1. 快速开始", 1)
    add_body(doc, "系统根据飞书账号识别用户，并根据管理员配置显示可访问的菜单。首次使用前，请确认账号已经同步到组织架构。")
    for text in ("在飞书内打开系统应用或系统地址。", "在登录页点击“使用飞书登录”，完成授权后进入工作台。", "根据需要进入“我的工时”“工时审批”或项目管理相关菜单。", "如果看不到菜单，请联系系统管理员检查模块权限和数据范围。"):
        add_step(doc, text)

    add_heading(doc, "2. 角色与菜单", 1)
    add_body(doc, "系统通过管理员管理配置功能权限和项目数据范围。没有权限的菜单不会显示，接口也不会返回超出范围的数据。")
    add_table(doc, ["角色", "常用操作", "数据范围"], [
        ("员工", "填报工时、查看本人记录和审批进度", "本人提交的工时"),
        ("工时审批人", "查看工时详情、通过或驳回", "分配给本人的审批任务"),
        ("项目负责人", "查看负责项目、合同阶段、经营提醒", "本人负责或授权项目"),
        ("出纳", "审批合同收款单和合同付款单", "分配给本人的收付款审批任务"),
        ("项目/业务管理员", "维护项目、合同、客户、供应商和单据", "按管理员配置的项目范围"),
        ("系统管理员", "组织同步、权限、基础配置和全局维护", "按系统管理员配置"),
    ], [1800, 3600, 3960])

    add_heading(doc, "3. 工作台", 1)
    add_body(doc, "工作台分为“我的项目”和“我的工时”两部分。项目负责人区展示负责项目的预算、人天、成本、收付款和风险；员工区展示个人工时和审批进度。")
    add_bullet(doc, "“填报工时”打开新增工时页面。")
    add_bullet(doc, "“我的项目”中的项目行可点击进入项目详情。")
    add_bullet(doc, "“项目提醒”显示合同阶段到期、逾期和预算风险。")
    add_bullet(doc, "“最近工时”点击具体记录进入该条工时详情，而不是仅进入列表。")
    add_note(doc, "数据口径", "项目经营数字按已审批工时和已审批收付款计算；待审批数据会单独显示，不会提前计入已收、已付或实际成本。", fill=LIGHT_GOLD, accent="A16207")

    add_heading(doc, "4. 工时填报", 1)
    add_heading(doc, "4.1 新增工时", 2)
    for text in ("进入“我的工时”，点击“新增工时”，也可以在工作台点击“填报工时”。", "先选择项目类型，再按项目名称或编号筛选项目。项目选择框只显示所选类型下的项目。", "填写开始日期、结束日期、工作阶段、工作状态和工作内容。当前工时单位为“天”。", "提交前检查项目、日期和工作内容，点击“保存并提交”提交本条工时。"):
        add_step(doc, text)
    add_screenshot(doc, "manual-worklog-form.png", "图 2  工时填报表单示例：先选项目类型，再选择项目", "工时填报表单界面示例")

    add_heading(doc, "4.2 填报校验规则", 2)
    add_table(doc, ["校验项", "规则"], [
        ("日期范围", "开始日期不能晚于结束日期，且开始和结束日期不能跨自然月。"),
        ("重复填报", "同一人员同一天不能填报多次，不同项目也不允许。"),
        ("区间重叠", "新日期区间不能与本人已有工时区间重叠。"),
        ("项目状态", "项目必须允许工时填报。"),
        ("工作内容", "应填写实际完成内容，不能留空。"),
    ], [2200, 7160])

    add_heading(doc, "4.3 查看和修改", 2)
    add_bullet(doc, "在“我的工时”列表点击记录，可查看完整信息和审批时间线。")
    add_bullet(doc, "草稿或已驳回记录可以修改后重新提交。")
    add_bullet(doc, "审批中和已通过工时不能由普通用户直接编辑；需要更正时联系工时管理员。")
    add_bullet(doc, "审批状态支持局部刷新和定时刷新，也可以手动点击刷新。")

    add_heading(doc, "5. 工时审批", 1)
    for text in ("进入“工时管理 > 工时审批”。", "在“待我审批”页签选择需要处理的工时。", "点击记录进入详情，核对人员、项目、日期、天数、工作阶段和工作描述。", "确认无误点击“通过”；需要修改时点击“驳回”并填写原因。"):
        add_step(doc, text)
    add_bullet(doc, "飞书卡片只负责通知和打开系统详情，审批操作在系统页面完成。")
    add_bullet(doc, "审批通过后，系统更新已发生人天、人员成本、我的工时和工时看板。")
    add_bullet(doc, "“已审批”页签可查看已通过和已驳回记录。")

    add_heading(doc, "6. 项目管理", 1)
    add_heading(doc, "6.1 项目列表", 2)
    add_body(doc, "进入“项目管理 > 项目管理”，可以按项目类型、项目状态、项目部门、项目负责人、客户以及项目名称或编号进行查询。部门筛选支持层级关系，选择上级部门时包含下级部门。")
    add_bullet(doc, "点击项目名称或整行进入项目详情。")
    add_bullet(doc, "新增和编辑使用弹窗，保存后项目列表自动刷新。")
    add_bullet(doc, "项目列表支持分页，预算和经营数字在项目详情中查看。")
    add_screenshot(doc, "manual-projects.png", "图 3  项目管理界面示例：筛选、列表与分页", "项目管理界面示例")

    add_heading(doc, "6.2 项目详情", 2)
    add_table(doc, ["区域", "主要内容"], [
        ("概况", "预算人天、已发生人天、剩余人天、成本预算、人员成本、合同成本和预警。"),
        ("基本信息", "名称、编号、类型、状态、负责人、部门、客户和是否允许填报工时。"),
        ("人员与工时", "工时审批人、上级领导和项目工时明细。"),
        ("预算", "人天预算和成本预算；已发生数据只读。"),
        ("合同与资金", "收入合同、支出合同、合同收款单和合同付款单。"),
        ("操作记录", "项目、预算、合同和收付款相关操作日志。"),
    ], [2200, 7160])
    add_note(doc, "组织架构来源", "项目负责人、审批人、上级领导和项目部门从飞书组织架构读取，系统不创建组织架构外的人员。", fill=LIGHT_BLUE)

    add_heading(doc, "7. 合同与收付款", 1)
    add_heading(doc, "7.1 合同和合同阶段", 2)
    for text in ("进入“项目管理 > 项目合同”，使用“收入合同/支出合同”页签切换类型。", "点击合同名称进入合同详情，合同详情默认展示合同信息和全部阶段。", "点击“查看阶段”只查看阶段列表；点击“新增阶段”才打开新增阶段表单。", "阶段的已收、已付和剩余金额由审批通过的合同收付款单自动计算。"):
        add_step(doc, text)
    add_screenshot(doc, "manual-contract-detail.png", "图 4  合同详情界面示例：合同信息与阶段明细平铺展示", "合同详情界面示例")
    add_body(doc, "实施项目可以同时建立收入合同和支出合同。研发、课题、观摩项目是否开放收入侧由项目类型配置决定。")

    add_heading(doc, "7.2 合同收款单", 2)
    for text in ("进入“合同收款”，点击新增。", "选择项目、收入合同和有剩余应收的合同阶段。", "填写本次金额、收款日期、付款方姓名、付款方开户行和付款方账号。", "确认金额不超过阶段剩余应收，提交出纳审批。"):
        add_step(doc, text)

    add_heading(doc, "7.3 合同付款单", 2)
    for text in ("进入“合同付款”，点击新增。", "选择项目、支出合同和有剩余应付的合同阶段。", "填写本次金额、付款日期、收款方姓名、收款方开户行和收款方账号。", "确认金额不超过阶段剩余应付，提交出纳审批。"):
        add_step(doc, text)
    add_note(doc, "必填提醒", "收款方/付款方姓名、开户行和银行账号是必填信息。供应商默认账户带出后，提交前仍需核对本单账户。", fill=LIGHT_RED, accent="B91C1C")

    add_heading(doc, "7.4 出纳审批", 2)
    for text in ("进入收款或付款列表的“待我审批”页签。", "点击单据名称或查看操作进入详情。", "核对项目、合同、阶段、金额、账户和备注。", "点击“通过”或“驳回”；驳回必须填写原因。"):
        add_step(doc, text)
    add_bullet(doc, "收款审批通过：增加已收、减少应收并计入项目收入。")
    add_bullet(doc, "付款审批通过：增加已付、减少应付并计入项目外部成本。")
    add_bullet(doc, "审批通过后系统向申请人发送飞书审批通过通知。")
    add_bullet(doc, "已审批单据不能直接编辑或删除，错误需要走冲销流程。")

    add_heading(doc, "8. 工时看板与项目看板", 1)
    add_heading(doc, "8.1 工时看板", 2)
    add_body(doc, "进入“工时管理 > 工时看板”，选择日期范围和部门后点击“查询”。工时看板只展示工时指标，不混入项目成本、预算和收付款。")
    add_table(doc, ["内容", "说明"], [
        ("核心指标", "已审批工时天数、填报人数、项目数、待审批数量和审批通过率。"),
        ("趋势和排行", "工时趋势、一级部门对比、项目排行和人员排行。"),
        ("分布图", "项目类型、工作阶段和审批状态分布。"),
        ("导出", "按当前筛选条件导出工时明细。"),
    ], [2200, 7160])

    add_heading(doc, "8.2 项目看板", 2)
    add_body(doc, "项目看板可按项目类型、状态、部门、负责人、客户和项目筛选，查看收入、总成本、应收、应付、已收、已付、毛利、毛利率、预算执行和风险预警。")
    add_screenshot(doc, "manual-project-board.png", "图 5  项目看板界面示例：经营汇总、趋势和账龄分析", "项目看板界面示例")

    add_heading(doc, "9. 系统管理", 1)
    add_heading(doc, "9.1 人员成本", 2)
    add_body(doc, "进入“系统管理 > 人员成本”，按部门和人员筛选后维护单天成本。新审批通过的工时按当前人员成本计算；已经生成的历史成本不会因为后续维护而自动重算。")
    add_heading(doc, "9.2 集成与通知", 2)
    add_body(doc, "集成与通知用于查看飞书连接状态、手动同步组织架构、查看同步结果和检查应用消息通知配置。组织同步后，人员、部门和项目人员字段使用最新数据。")
    add_heading(doc, "9.3 管理员管理", 2)
    add_body(doc, "管理员管理用于为指定人员配置模块权限和项目数据范围，包括本人负责项目、指定部门及下级部门、全部项目。未配置的模块不会显示。")
    add_heading(doc, "9.4 基础配置", 2)
    add_body(doc, "基础配置用于维护项目类型、项目状态、工时类型、工作状态等下拉选项。选项支持新增、编辑、排序、启用和停用；已经被业务数据引用的选项不能删除。")

    add_heading(doc, "10. 常见问题与安全提示", 1)
    add_heading(doc, "10.1 常见问题", 2)
    add_table(doc, ["问题", "处理方式"], [
        ("看不到项目或人员", "确认组织架构已同步，并检查项目类型、项目状态、项目是否允许填报和权限范围。"),
        ("工时提交失败", "检查日期是否跨月、是否同一天重复、日期区间是否重叠、项目是否允许填报。"),
        ("审批状态没有立即变化", "使用页面刷新；系统支持局部刷新和定时刷新。"),
        ("没有收到飞书通知", "管理员检查应用发布、消息权限、组织同步和集成与通知页面状态。"),
        ("合同收付款单无法提交", "确认合同类型、阶段剩余金额、姓名、开户行和银行账号均正确填写。"),
        ("看不到菜单", "联系系统管理员配置对应模块权限和数据范围。"),
    ], [2600, 6760])
    add_heading(doc, "10.2 安全提示", 2)
    add_bullet(doc, "不要在聊天、文档或截图中发送飞书 App Secret、Cloudflare token、SSH 密码或完整银行账号。")
    add_bullet(doc, "银行账号在列表和飞书卡片中默认脱敏，只在有权限的单据详情中查看。")
    add_bullet(doc, "不要手工改写已经审批通过的金额；发现错误应使用冲销流程。")
    add_bullet(doc, "工时看板默认按已审批工时统计，待审批和驳回数据不会混入已审批口径。")
    add_note(doc, "支持入口", "遇到系统异常时，请提供页面名称、操作步骤、时间和提示信息，联系系统管理员；不要提供密码、密钥或完整银行账号。", fill=LIGHT_BLUE)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()

