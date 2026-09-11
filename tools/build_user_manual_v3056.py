from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "assets" / "manual-v3.0.56"
OUT_DOCX = DOCS / "综合项目管理系统用户使用手册-v3.0.56.docx"
OUT_MD = DOCS / "01-user-manual-v3.0.56.md"
CANONICAL_MD = DOCS / "01-user-manual.md"

FONT = "Microsoft YaHei"
TITLE = "综合项目管理系统"
VERSION = "v3.0.56"
UPDATE_DATE = "2026-09-07"

TEAL = "0F766E"
TEAL_DARK = "134E4A"
BLACK = "000000"
SLATE = "243B3A"
MUTED = "64748B"
LIGHT_TEAL = "E6FFFA"
LIGHT_SLATE = "F8FAFC"
LIGHT_AMBER = "FFFBEB"
LIGHT_RED = "FEF2F2"
BORDER = "CBD5E1"


def set_font(run, size: float = 10.5, color: str = SLATE, bold: bool | None = None, italic: bool | None = None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    rfonts.set(qn("w:eastAsia"), FONT)


def set_paragraph_spacing(paragraph, before: float = 0, after: float = 6, line: float = 1.2):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def set_cell_border(cell, color: str = BORDER, size: str = "6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        edge = borders.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), size)
        edge.set(qn("w:color"), color)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.first_child_found_in("w:shd")
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def paragraph(doc: Document, text: str = "", *, size: float = 10.5, color: str = SLATE, bold: bool = False,
              italic: bool = False, style: str | None = None, after: float = 6):
    p = doc.add_paragraph(style=style)
    set_paragraph_spacing(p, after=after)
    if text:
        r = p.add_run(text)
        set_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading("", level=level)
    set_paragraph_spacing(p, before=8 if level == 1 else 4, after=6)
    r = p.add_run(text)
    set_font(r, size=18 if level == 1 else 13 if level == 2 else 11.5, color=BLACK, bold=True)
    return p


def bullet(doc: Document, text: str, level: int = 0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    set_paragraph_spacing(p, after=3)
    r = p.add_run(text)
    set_font(r, size=10, color=SLATE)
    return p


def step(doc: Document, text: str):
    p = doc.add_paragraph(style="List Number")
    set_paragraph_spacing(p, after=3)
    r = p.add_run(text)
    set_font(r, size=10.3, color=SLATE)
    return p


def note(doc: Document, title: str, body: str, fill: str = LIGHT_TEAL, accent: str = TEAL):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=6)
    r1 = p.add_run(title)
    set_font(r1, size=10.5, color=BLACK, bold=True)
    r2 = p.add_run("  " + body)
    set_font(r2, size=10, color=SLATE)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    hrow = table.rows[0]
    set_repeat_header(hrow)
    for idx, text in enumerate(headers):
        cell = hrow.cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "ECFDF5")
        set_cell_border(cell)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_font(r, size=9.5, color=TEAL_DARK, bold=True)
        if widths:
            cell.width = Inches(widths[idx])
    for values in rows:
        row = table.add_row()
        for idx, text in enumerate(values):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            set_paragraph_spacing(p, after=0)
            r = p.add_run(str(text))
            set_font(r, size=9.2, color=SLATE)
            if widths:
                cell.width = Inches(widths[idx])
    paragraph(doc, "", after=2)
    return table


def add_screenshot(doc: Document, filename: str, caption: str, alt: str):
    path = ASSETS / filename
    if not path.exists():
        note(doc, "截图缺失", f"{filename} 未找到，手册生成时跳过该图。", fill=LIGHT_AMBER, accent="B45309")
        return
    with Image.open(path) as img:
        width, height = img.size
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=4, after=2)
    available_width = 6.65
    # Tall screenshots use a slightly smaller width to keep them on a readable page.
    insert_width = 6.35 if height / width > 0.75 else available_width
    shape = p.add_run().add_picture(str(path), width=Inches(insert_width))
    shape._inline.docPr.set("descr", alt)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(cap, after=8)
    r = cap.add_run(caption)
    set_font(r, size=8.8, color=MUTED, italic=True)


def add_page_number(section):
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("综合项目管理系统用户使用手册 · ")
    set_font(run, size=8.5, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def configure_doc(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)
    add_page_number(section)
    styles = doc.styles
    for style_name in ("Normal", "List Bullet", "List Bullet 2", "List Number"):
        style = styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(10.5)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def cover(doc: Document):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=90, after=4)
    r = p.add_run(TITLE)
    set_font(r, size=28, color=BLACK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=20)
    r = p.add_run("用户使用手册")
    set_font(r, size=22, color=BLACK, bold=True)

    for text in [
        f"文档版本：{VERSION}",
        f"更新时间：{UPDATE_DATE}",
        "适用对象：员工、审批人、项目负责人、出纳、财务/经营审批、系统管理员",
        "使用入口：飞书应用内入口或管理员提供的系统地址",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=5)
        r = p.add_run(text)
        set_font(r, size=11, color=MUTED)

    note(
        doc,
        "阅读提示",
        "本手册按实际操作路径编写。普通员工重点看工作台、我的工时和项目费用；项目负责人重点看工时审批、项目详情和看板；管理员重点看系统管理章节。",
    )
    doc.add_page_break()


def add_revision_summary(doc: Document):
    heading(doc, "0. 本版手册说明", 1)
    paragraph(doc, "本版手册基于 2026-09-07 已发布的 v3.0.56 版本重新生成，并按此前要求补充了更细的图文步骤。")
    add_table(
        doc,
        ["改进点", "本版处理方式"],
        [
            ["系统名称统一", "全文统一为“综合项目管理系统”，英文标识使用“SINYE PROJECT MANAGE SYSTEM”。"],
            ["截图比例", "重新截取 1250×700 左右常规浏览器视窗，插入 Word 时统一宽度，避免 4K 截图缩在左上角。"],
            ["操作颗粒度", "每个核心模块均给出进入路径、操作步骤、结果说明和常见注意事项。"],
            ["最新功能范围", "补齐项目费用、动态审批、通知中心、角色权限、附件上传预览、项目经营分析等新功能。"],
        ],
        widths=[1.8, 4.6],
    )
    note(doc, "敏感信息", "手册不记录服务器密码、飞书应用密钥等敏感配置。生产地址和账号权限以管理员实际配置为准。", fill=LIGHT_AMBER, accent="B45309")


def add_contents(doc: Document):
    heading(doc, "1. 目录与角色速查", 1)
    add_table(
        doc,
        ["角色", "建议优先阅读", "典型工作"],
        [
            ["普通员工", "工作台、我的工时、项目费用、常见问题", "填报工时、提交费用/报销、查看审批状态。"],
            ["工时审批人", "工时审批、工时详情、飞书通知", "收到飞书提醒后进入系统审批工时。"],
            ["项目负责人", "工作台项目区、项目详情、合同与资金、项目看板", "查看项目预算、成本、合同阶段、收付款和预警。"],
            ["财务/出纳", "合同收款单、合同付款单、费用审批、通知中心", "处理收付款确认、出纳支付和相关审批。"],
            ["系统管理员", "系统管理、权限、审批岗位、基础配置、集成与通知", "维护组织同步、人员成本、角色权限、审批流程和下拉配置。"],
        ],
        widths=[1.3, 2.2, 2.9],
    )
    paragraph(doc, "手册主线如下：登录进入系统 → 在工作台查看待办 → 员工填报工时或提交费用 → 审批人处理审批 → 项目负责人查看项目经营和预警 → 管理员维护权限、岗位与基础配置。")


def add_login_and_workbench(doc: Document):
    heading(doc, "2. 登录、首页与菜单", 1)
    heading(doc, "2.1 登录系统", 2)
    step(doc, "在飞书内打开系统应用，或使用管理员提供的系统地址访问。")
    step(doc, "点击“使用飞书登录”。如系统配置为飞书连接模式，会向飞书申请当前用户身份。")
    step(doc, "首次接入时，系统会将第一个成功接入的飞书用户初始化为系统管理员；后续人员权限由管理员维护。")
    step(doc, "登录后系统根据角色、菜单权限、按钮权限和数据范围显示可用功能。没有权限的菜单不会展示，或进入后只能看到授权范围内的数据。")
    note(doc, "登录失败排查", "若提示 FEISHU_USER_NOT_SYNCED，通常是飞书用户尚未同步到系统；请管理员在“集成与通知”中同步组织架构，或检查飞书应用通讯录权限。", fill=LIGHT_AMBER, accent="B45309")

    heading(doc, "2.2 工作台总览", 2)
    paragraph(doc, "工作台分为上、下两块：上半部分服务项目负责人，展示负责项目的经营概况和提醒；下半部分服务员工，展示本月工时、待办和最近工时。")
    add_screenshot(doc, "01-workbench.png", "图 2-1 工作台：上半区为项目负责人视角，下半区为员工视角。", "工作台首页截图")
    step(doc, "查看“我的项目”区，确认负责项目数量、待审工时、7 天内应收/应付、逾期事项和预警项目。")
    step(doc, "在“项目经营概况”中点击项目记录，可以进入项目详情查看预算、合同、费用和操作记录。")
    step(doc, "在“我的工时”区点击“填报工时”发起工时；点击最近工时记录可直接进入该条工时详情。")
    step(doc, "欢迎语会按 Asia/Shanghai 当前时间自动变化，例如早上好、上午好、中午好、下午好或晚上好。日期行只显示日期和星期。")

    heading(doc, "2.3 菜单分组", 2)
    add_table(
        doc,
        ["大模块", "菜单", "用途"],
        [
            ["工作台", "工作台", "查看个人待办、负责项目概况、快速发起工时/收付款/费用。"],
            ["工时管理", "我的工时、工时审批、工时看板、工时台账", "完成工时填报、审批、统计和后台工时维护。"],
            ["项目管理", "项目管理、项目合同、合同收款、合同付款、项目费用、项目看板、供应商、客户", "维护项目主数据、合同、阶段、收付款、费用和项目经营分析。"],
            ["系统管理", "人员成本、集成与通知、管理员管理、审批岗位、审批流程、基础配置", "维护组织、权限、成本、流程、通知策略和基础下拉项。"],
        ],
        widths=[1.2, 2.4, 3.0],
    )


def add_worklog(doc: Document):
    heading(doc, "3. 工时管理", 1)
    heading(doc, "3.1 我的工时", 2)
    paragraph(doc, "“我的工时”用于员工查看本人所有工时，包含草稿、审批中、已通过、已驳回等状态。")
    add_screenshot(doc, "02-worklogs.png", "图 3-1 我的工时列表：可查看状态、审批人和每条工时的操作入口。", "我的工时列表截图")
    step(doc, "进入“工时管理 > 我的工时”。")
    step(doc, "通过状态筛选查看草稿、审批中、已驳回或已通过记录。")
    step(doc, "点击“查看详情”查看单条工时完整信息；草稿或驳回记录点击“编辑”后重新提交。")
    step(doc, "列表状态会自动刷新；如刚被审批但状态未即时变化，可稍等片刻或刷新当前页面。")

    heading(doc, "3.2 新增并提交工时", 2)
    add_screenshot(doc, "03-worklog-form.png", "图 3-2 新增工时：先选项目类型，再选择同类型项目。", "新增工时页面截图")
    step(doc, "点击“＋ 新增工时”。")
    step(doc, "先选择项目类型，例如研发项目、课题项目、实施项目或观摩项目。")
    step(doc, "在项目选择框中按名称筛选项目。选定项目类型后，只能选择该类型下且允许填报的项目。")
    step(doc, "填写开始日期和结束日期。系统自动按工作日计算工时天数，单位为“天”。")
    step(doc, "填写项目阶段、工时类型、工作状态和工作描述。工作描述应说明实际完成内容，便于审批人判断。")
    step(doc, "点击“保存草稿”可暂存；点击“提交审批”会生成单条工时审批任务并发送飞书通知。")
    note(doc, "工时校验规则", "同一人员同一天不允许重复填报，即使不同项目也不允许；日期区间不能跨自然月；开始日期不能晚于结束日期；工时天数必须为整数工作日。", fill=LIGHT_AMBER, accent="B45309")

    heading(doc, "3.3 查看工时详情", 2)
    add_screenshot(doc, "05-worklog-detail.png", "图 3-3 工时详情：展示单号、状态、工作内容、通知结果和审批时间线。", "工时详情截图")
    step(doc, "在“我的工时”点击某条记录的“查看详情”。")
    step(doc, "核对工时单号、项目、阶段、日期范围、天数和描述。")
    step(doc, "在右侧审批时间线查看提交、确认审批人、发送飞书通知、审批通过或驳回等节点。")
    step(doc, "若为审批中状态，可根据通知状态判断飞书提醒是否已发送；业务状态以系统本地审批记录为准。")

    heading(doc, "3.4 工时审批", 2)
    paragraph(doc, "“工时审批”分为“待我审批”和“已审批”两个页签。飞书工时审批通知卡片会直达该条工时审批详情页，审批动作仍在系统内完成。")
    add_screenshot(doc, "04-worklog-approvals.png", "图 3-4 工时审批：审批人查看待办和已处理记录。", "工时审批页面截图")
    step(doc, "进入“工时管理 > 工时审批”。")
    step(doc, "在“待我审批”中点击单据详情，查看人员、项目、日期、天数和工作说明。")
    step(doc, "确认无误点击“通过”；需要员工补充或更正时点击“驳回”并填写原因。")
    step(doc, "审批通过后，系统更新工时状态、项目已发生人天，并按审批时人员单天成本生成项目人员成本。")
    note(doc, "自审转上级", "如果项目审批人配置为提交人本人，系统会按项目配置转交上级领导审批，避免自己审批自己的工时。")

    heading(doc, "3.5 工时看板", 2)
    paragraph(doc, "工时看板只保留工时分析，不混入项目成本预算指标。适合查看部门、项目、人员、日期趋势和审批状态。")
    add_screenshot(doc, "06-worklog-board.png", "图 3-5 工时看板：按部门、日期范围查询工时趋势和分布。", "工时看板截图")
    step(doc, "进入“工时管理 > 工时看板”。")
    step(doc, "选择日期范围和部门。部门筛选支持组织层级，选择上级部门会包含下级部门。")
    step(doc, "点击“查询”刷新看板。")
    step(doc, "需要明细时点击“导出明细”。如果在飞书内置浏览器没有弹出下载窗口，请检查下载栏或浏览器下载记录。")

    heading(doc, "3.6 工时台账", 2)
    paragraph(doc, "工时台账是管理员后台维护功能，用于查看所有人的工时日志，支持按日期、部门、人员和项目筛选，并可按权限进行修改或删除。")
    add_screenshot(doc, "07-worklog-ledger.png", "图 3-6 工时台账：后台按条件检索所有工时并分页维护。", "工时台账截图")
    step(doc, "进入“工时管理 > 工时台账”。")
    step(doc, "使用日期、部门、人员、项目筛选条件定位记录。")
    step(doc, "根据每页 10/20/50 条或自定义条数分页查看。")
    step(doc, "仅有相应按钮权限的管理员可以修改或删除记录。修改会保留审计记录。")


def add_project(doc: Document):
    heading(doc, "4. 项目管理", 1)
    heading(doc, "4.1 项目列表", 2)
    paragraph(doc, "项目管理主页面只保留项目主信息，预算、收入、成本、应收应付等经营数据在项目详情和项目看板查看。")
    add_screenshot(doc, "08-projects.png", "图 4-1 项目管理：筛选项紧凑排列，列表展示项目主信息。", "项目管理列表截图")
    step(doc, "进入“项目管理 > 项目管理”。")
    step(doc, "可按项目类型、项目部门、项目状态、项目负责人、客户和项目标题模糊查询。")
    step(doc, "部门筛选是树形层级，选择上级部门时会包含全部下级部门。")
    step(doc, "项目列表支持分页；点击“查看”进入项目详情，点击“编辑”打开项目信息弹窗。")

    heading(doc, "4.2 新增或编辑项目", 2)
    add_screenshot(doc, "09-project-form.png", "图 4-2 项目表单：在弹窗内维护主信息、预算和审批配置。", "项目新增编辑弹窗截图")
    step(doc, "点击“＋ 新增项目”或列表中的“编辑”。")
    step(doc, "填写项目编码、项目名称、项目状态、项目类型、项目负责人、所属部门和客户。")
    step(doc, "设置工时填报开关。关闭后员工不能继续向该项目提交工时。")
    step(doc, "维护人天预算、成本预算、费用预算等预算字段。")
    step(doc, "配置工时审批人、上级领导等审批相关人员。所有人员与部门都来自飞书组织架构。")
    step(doc, "点击“保存项目”。弹窗不会因为误点遮罩自动消失，减少未保存内容丢失。")

    heading(doc, "4.3 项目详情", 2)
    paragraph(doc, "项目详情用于单一项目的细致分析，分为概况、基本信息、人员与工时、预算、合同与资金、费用明细、操作记录等页签。")
    add_screenshot(doc, "10-project-detail.png", "图 4-3 项目详情：顶部展示收入合同、支出合同、总成本、项目费用、毛利和预算风险。", "项目详情截图")
    step(doc, "在项目列表点击“查看”。")
    step(doc, "在顶部 KPI 区查看收入合同、支出合同、总成本、项目费用、毛利/毛利率和未完成预算。")
    step(doc, "在“成本组成与预算执行”查看人员成本、合同成本、项目费用和预算剩余。")
    step(doc, "切换“合同与资金”查看收入合同、支出合同、收款、付款和阶段余额。")
    step(doc, "切换“费用明细”查看差旅、无合同发票和其他费用单据。")
    step(doc, "切换“操作记录”追踪项目、预算、合同、收付款、费用等变更。")
    note(doc, "成本口径", "项目实际总成本用于经营分析，包含已确认人员成本、已确认合同付款成本和已审批项目费用成本。已收、已付体现资金确认结果；应收、应付按到期阶段余额计算。")


def add_contract_and_cash(doc: Document):
    heading(doc, "5. 合同、阶段与收付款", 1)
    heading(doc, "5.1 项目合同列表", 2)
    paragraph(doc, "合同分为收入合同和支出合同。实施项目既可能有收入合同，也可能有支付/支出合同。")
    add_screenshot(doc, "11-contracts.png", "图 5-1 项目合同：收入、支出合同分开维护，并支持导入。", "项目合同列表截图")
    step(doc, "进入“项目管理 > 项目合同”。")
    step(doc, "使用收入合同/支出合同页签切换类型。")
    step(doc, "可按项目、合同类型、往来单位、合同状态和日期查询。")
    step(doc, "点击合同名称进入合同详情；点击“查看阶段”只查看阶段列表；点击“编辑”修改合同主信息。")

    heading(doc, "5.2 新增合同与保存草稿", 2)
    add_screenshot(doc, "12-contract-form.png", "图 5-2 合同表单：新建时即可上传附件，支持保存草稿。", "合同新增编辑弹窗截图")
    step(doc, "点击“＋ 新增合同”。")
    step(doc, "选择项目、合同类型和往来单位。收入合同选择客户，支出合同选择供应商。")
    step(doc, "填写合同名称、合同编号、合同金额、签订日期、生效日期等核心字段。")
    step(doc, "按需上传合同附件。附件在新建时即可上传，不需要先保存合同。")
    step(doc, "点击“保存草稿”可先落库但不进入执行；确认无误后保存为执行中或后续激活。")
    note(doc, "合同导入", "导入合同时，项目字段填写项目名称；合同类型填写中文“收入合同/支出合同”；客户或供应商填写名称，系统匹配后自动保存对应编码。匹配失败不会写入，需要修正模板后重新导入。")

    heading(doc, "5.3 合同详情和阶段", 2)
    add_screenshot(doc, "13-contract-detail.png", "图 5-3 合同详情：合同主信息、结算率、阶段和款单笔数集中展示。", "合同详情截图")
    step(doc, "从合同列表点击合同名称进入详情。")
    step(doc, "查看合同金额、状态、项目、客户/供应商、结算率、已收/已付和剩余金额。")
    step(doc, "在阶段区域查看每个阶段的计划金额、计划日期、已收/已付、剩余金额和款单笔数。")
    step(doc, "点击“新增阶段”才打开新增阶段弹窗；普通点击阶段只查看阶段列表，避免误以为正在新增。")
    step(doc, "附件区域位于合同详情底部，可上传、下载或预览。Word、Excel、PPT 等 Office 文件会尽量转为可预览格式；转换失败时可下载本地查看。")

    heading(doc, "5.4 合同收款单", 2)
    add_screenshot(doc, "14-receipts.png", "图 5-4 合同收款：查看待审批和全部合同收款单。", "合同收款列表截图")
    add_screenshot(doc, "15-receipt-form.png", "图 5-5 新增合同收款单：选择收入合同后带出客户付款方信息。", "新增合同收款单截图")
    step(doc, "进入“项目管理 > 合同收款”，点击“＋ 新增合同收款单”。")
    step(doc, "选择项目、收入合同和合同阶段。")
    step(doc, "系统会根据合同关联客户自动带出付款方名称、付款银行和付款账号；提交前仍可按本单实际情况修改。")
    step(doc, "填写本次收款金额、收款日期、备注，并上传收款附件。")
    step(doc, "提交后进入出纳审批或确认流程；审批通过后通知申请人，增加项目已收并核减已到期应收。")

    heading(doc, "5.5 合同付款单", 2)
    add_screenshot(doc, "16-payments.png", "图 5-6 合同付款：合同付款统一纳入动态审批流程。", "合同付款列表截图")
    add_screenshot(doc, "17-payment-form.png", "图 5-7 新增合同付款单：选择支出合同后带出供应商收款账户。", "新增合同付款单截图")
    step(doc, "进入“项目管理 > 合同付款”，点击“＋ 新增合同付款单”。")
    step(doc, "选择项目、支出合同和合同阶段。")
    step(doc, "系统会根据供应商默认账户自动带出收款方名称、收款银行和收款账号；账号、收款方姓名和银行为必填。")
    step(doc, "填写本次付款金额、计划付款日期、备注，并上传付款附件。")
    step(doc, "合同付款、费用付款和报销款使用同一类付款审批路径。出纳支付节点完成后，系统增加项目已付并核减应付。")
    note(doc, "应收/应付规则", "项目看板中的“应收未收”只统计已到期收入阶段的未收金额；“应付未付”只统计已到期支出阶段的未付金额。未到期阶段用于临期提醒，不计入已到期未收/未付。")


def add_expenses(doc: Document):
    heading(doc, "6. 项目费用", 1)
    paragraph(doc, "项目费用覆盖差旅、无合同发票和其他费用，用于补充合同之外的项目成本。费用审批通过后计入项目费用成本，出纳支付后计入已付资金。")
    add_screenshot(doc, "18-project-expenses.png", "图 6-1 项目费用：费用台账、待审批和已处理记录。", "项目费用列表截图")
    add_screenshot(doc, "19-expense-form.png", "图 6-2 新增费用：新建时即可上传附件并提交审批。", "新增项目费用截图")
    heading(doc, "6.1 提交项目费用", 2)
    step(doc, "进入“项目管理 > 项目费用”，点击“＋ 新增费用”。")
    step(doc, "选择费用类型：差旅费用、无合同发票或其他费用。")
    step(doc, "选择所属项目，填写费用金额、费用日期、事由和说明。")
    step(doc, "上传发票、行程单、合同外付款凭证等附件。附件可在新建时直接上传。")
    step(doc, "点击保存草稿或提交审批。提交前系统会校验审批岗位是否配置齐全。")
    heading(doc, "6.2 费用审批流程", 2)
    add_table(
        doc,
        ["费用类型", "审批路径"],
        [
            ["差旅费用", "经办人提交 → 项目负责人审批 → 会计审批 → 上级领导审批 → 财务总监审批 → 总经理审批 → 出纳支付。"],
            ["费用付款 / 无合同发票 / 其他费用 / 合同付款", "经办人提交 → 项目负责人审批 → 会计审批 → 经营审批 → 上级领导审批 → 经营副总审批 → 财务总监审批 → 总经理审批 → 出纳支付。"],
        ],
        widths=[1.8, 4.9],
    )
    step(doc, "当前节点的全部候选审批人会收到飞书应用消息。")
    step(doc, "任一候选人处理后，审批记录显示实际审批人姓名。")
    step(doc, "流程进入下一节点时，系统继续向下一节点候选人发送飞书提醒。")
    step(doc, "审批记录中的动作显示为中文，例如“提交”“审批通过”“驳回”“出纳支付”。")
    note(doc, "流程设计原则", "审批人员不写死在代码里，由“审批岗位”和“审批流程”配置解析；项目负责人、上级领导等节点会结合项目配置和飞书组织关系动态确定。")


def add_boards(doc: Document):
    heading(doc, "7. 项目看板与经营预警", 1)
    paragraph(doc, "项目看板用于项目经营视角，不替代财务总账。它把项目收入、成本、预算、阶段收付款和预警集中展示，支持项目负责人查看自己项目，具备全量权限的用户查看全部项目。")
    add_screenshot(doc, "21-project-board.png", "图 7-1 项目看板：经营汇总、项目列表和可下钻风险。", "项目看板截图")
    heading(doc, "7.1 查询项目看板", 2)
    step(doc, "进入“项目管理 > 项目看板”。")
    step(doc, "选择统计范围：我的项目或全部项目。选择全部项目时，项目选择框会动态获取当前可访问的全部项目。")
    step(doc, "按项目类型、状态、部门、负责人、客户和项目筛选。")
    step(doc, "点击查询后查看经营汇总、预算执行、成本组成、收付款趋势和预警明细。")
    heading(doc, "7.2 看板核心指标", 2)
    add_table(
        doc,
        ["指标", "含义"],
        [
            ["收入", "收入合同及已确认收款相关数据，用于观察项目收入侧情况。"],
            ["总成本", "人员成本、合同成本和项目费用成本的合计，用于项目经营分析。"],
            ["人员成本", "审批通过工时 × 审批时人员单天成本。人员成本后维护时，系统可回算 0 成本快照。"],
            ["合同成本", "支出合同付款按审批/支付确认后的成本归集。"],
            ["项目费用", "差旅、无合同发票和其他费用审批通过后形成费用成本。"],
            ["应收未收", "只统计已到期收入阶段的未收金额。"],
            ["应付未付", "只统计已到期支出阶段的未付金额。"],
            ["毛利/毛利率", "按收入与实际总成本计算，用于项目经营判断。"],
        ],
        widths=[1.5, 5.0],
    )
    heading(doc, "7.3 系统提醒和飞书预警", 2)
    add_table(
        doc,
        ["提醒类型", "触发对象", "展示/通知位置"],
        [
            ["工时待审批", "工时提交后的项目审批人或上级领导", "工时审批列表、飞书卡片直达该工时详情。"],
            ["合同收款单待审批", "具备出纳/收款审批职责的人员", "合同收款、通知中心、飞书提醒。"],
            ["合同付款单待审批", "付款审批流程当前节点候选人", "合同付款、通知中心、飞书提醒。"],
            ["项目费用待审批", "费用流程当前节点候选人", "项目费用、通知中心、飞书提醒。"],
            ["合同阶段临期", "项目负责人", "工作台项目提醒、项目看板预警、飞书提醒。"],
            ["合同阶段到期/逾期", "项目负责人", "工作台项目提醒、项目看板预警、飞书提醒。"],
            ["成本预算接近或超支", "项目负责人", "项目看板预警、飞书提醒。"],
            ["人天预算接近或超支", "项目负责人", "项目看板预警、飞书提醒。"],
        ],
        widths=[1.8, 2.4, 2.5],
    )
    note(doc, "提醒窗口", "工作台项目提醒会展示比“未来 7 天统计”更宽的范围，可列出 7 天以后但一个月以内的阶段，便于提前处理收付款。")


def add_counterparties_and_cost(doc: Document):
    heading(doc, "8. 客户、供应商与人员成本", 1)
    heading(doc, "8.1 供应商管理", 2)
    add_screenshot(doc, "22-suppliers.png", "图 8-1 供应商管理：维护供应商核心信息和默认收款账户。", "供应商管理截图")
    step(doc, "进入“项目管理 > 供应商管理”。")
    step(doc, "新增供应商时只填写必要信息：供应商名称、联系人、电话、地址、税号或营业执照信息、收款方名称、开户行和账号。")
    step(doc, "供应商编码由系统自动生成，导入模板不需要填写供应商编码。")
    step(doc, "新建合同付款单选择支出合同后，系统会带出供应商账户信息。")

    heading(doc, "8.2 客户管理", 2)
    add_screenshot(doc, "23-customers.png", "图 8-2 客户管理：维护客户核心信息和默认付款账户。", "客户管理截图")
    step(doc, "进入“项目管理 > 客户管理”。")
    step(doc, "新增客户时填写客户名称、联系人、电话、地址、税号或营业执照信息、付款方名称、开户行和账号。")
    step(doc, "客户编码由系统自动生成。")
    step(doc, "新建合同收款单选择收入合同后，系统会带出客户付款账户信息。")

    heading(doc, "8.3 人员成本", 2)
    add_screenshot(doc, "24-personnel-costs.png", "图 8-3 人员成本：按部门和人员维护单天成本。", "人员成本管理截图")
    step(doc, "进入“系统管理 > 人员成本”。")
    step(doc, "使用部门树和人员筛选定位员工。选择上级部门时包含下级部门。")
    step(doc, "点击添加或编辑，维护人员单天成本。")
    step(doc, "保存后，新的已审批工时按该成本生成项目人员成本；历史 0 成本快照可由系统回算补齐。")
    note(doc, "成本使用率", "人员成本使用率采用已发生人员成本 ÷ 按已审批工时加权的人均单天成本预算基准。该口径用于估算人天预算对应的人员成本消耗。")


def add_system_admin(doc: Document):
    heading(doc, "9. 系统管理", 1)
    heading(doc, "9.1 集成与通知", 2)
    add_screenshot(doc, "25-integrations.png", "图 9-1 集成与通知：飞书连接、组织同步、通知策略统一维护。", "集成与通知截图")
    step(doc, "进入“系统管理 > 集成与通知”。")
    step(doc, "在顶部查看飞书连接状态、组织成员数、部门数、今日发送数和待重试数。")
    step(doc, "点击“同步组织架构”从飞书读取成员和部门，更新人员、部门、项目负责人、审批人等选择项。")
    step(doc, "通知策略平时为只读状态；点击右上角“修改”后在弹窗内统一编辑，再一次保存。")
    step(doc, "飞书通知只负责提醒和打开系统页面，业务审批状态以本系统数据库为准。")

    heading(doc, "9.2 通知中心", 2)
    add_screenshot(doc, "26-notification-center.png", "图 9-2 通知中心：查看发送状态、失败原因和重试入口。", "通知中心截图")
    step(doc, "在“集成与通知”中切换到“通知中心”。")
    step(doc, "按业务类型、接收人和发送状态查看通知记录。")
    step(doc, "失败记录可查看失败原因；具备权限的管理员可点击重试。")
    step(doc, "审批类飞书卡片会携带深链，审批人点击后进入对应工时、付款、收款或费用详情页。")

    heading(doc, "9.3 管理员管理与权限", 2)
    add_screenshot(doc, "27-admin-permissions.png", "图 9-3 管理员管理：角色、菜单、按钮和数据范围集中配置。", "管理员管理截图")
    step(doc, "进入“系统管理 > 管理员管理”。")
    step(doc, "只把需要管理权限的人员添加进来；不需要管理的人不显示在管理员列表。")
    step(doc, "为人员绑定角色。角色由菜单权限、按钮权限和数据访问范围组成。")
    step(doc, "菜单权限决定能否看到模块；按钮权限决定能否新增、编辑、删除、导入、导出、审批、重试等；数据范围决定能看本人、负责项目、本部门及下级或全部数据。")
    step(doc, "查询、重置、分页、TAB 切换、关闭弹窗、返回上一页等普通界面按钮不单独配置权限。")

    heading(doc, "9.4 审批岗位", 2)
    add_screenshot(doc, "28-approval-positions.png", "图 9-4 审批岗位：维护会计、经营审批、财务总监、总经理、出纳等岗位成员。", "审批岗位截图")
    step(doc, "进入“系统管理 > 审批岗位”。")
    step(doc, "为会计、经营审批、经营副总、财务总监、总经理、出纳等岗位添加成员。")
    step(doc, "凡是配置为当前审批节点候选人的人员，都会收到飞书待办通知。")
    step(doc, "某个候选人完成审批后，审批记录显示实际处理人的姓名。")

    heading(doc, "9.5 审批流程", 2)
    add_screenshot(doc, "29-approval-workflows.png", "图 9-5 审批流程：按单据类型展示流程节点。", "审批流程截图")
    step(doc, "进入“系统管理 > 审批流程”。")
    step(doc, "查看差旅费用流程和统一付款/费用流程的节点顺序。")
    step(doc, "流程模板按版本管理。已提交单据沿用提交时的流程版本，新流程只影响后续新单据。")
    step(doc, "提交单据前，如某个必经岗位未配置人员，系统会提示管理员先补齐岗位。")

    heading(doc, "9.6 基础配置", 2)
    add_screenshot(doc, "30-configuration.png", "图 9-6 基础配置：维护项目类型、状态、工作类型等动态下拉。", "基础配置截图")
    step(doc, "进入“系统管理 > 基础配置”。")
    step(doc, "维护项目类型、项目状态、执行状态、工时类型、项目阶段等下拉项。")
    step(doc, "项目类型当前保留：研发项目、课题项目、实施项目、观摩项目。")
    step(doc, "停用或删除不再使用的下拉项，避免新增项目和工时填报时出现多余选项。")


def add_faq_and_data(doc: Document):
    heading(doc, "10. 常见问题与数据口径", 1)
    add_table(
        doc,
        ["问题", "处理建议"],
        [
            ["看不到项目", "检查是否为项目负责人、是否被授予项目数据范围、项目是否在当前筛选条件内、所属部门是否同步正确。"],
            ["看不到菜单", "请系统管理员检查角色菜单权限。只有被授予的管理模块才会显示。"],
            ["工时提交失败", "检查是否同一天已填报、日期是否跨月、项目是否允许填报、天数是否为系统计算的整数工作日。"],
            ["飞书没有收到通知", "先在通知中心查看是否生成通知记录，再看发送状态、失败原因和接收人 open_id 是否同步。"],
            ["审批人系统里没有待审", "检查该单据当前流程节点、候选人岗位配置、数据权限是否拦截。被指派审批任务应允许审批人进入详情处理。"],
            ["导出后没有弹下载", "飞书内置浏览器可能不弹出下载窗口，请查看下载栏/下载记录，或换外部浏览器下载。"],
            ["附件预览失败", "Office 文件需要后台转换为预览格式；如果转换失败，可下载附件本地查看。"],
            ["合同或合同收付款单位名称变化", "系统应按关联 ID 动态显示客户/供应商最新名称。如页面旧值未刷新，请重新进入页面或点击查询。"],
        ],
        widths=[2.0, 4.8],
    )

    heading(doc, "10.1 关键计算口径", 2)
    add_table(
        doc,
        ["口径", "说明"],
        [
            ["工时天数", "按开始日期至结束日期包含首尾的工作日计算，只统计周一至周五，暂不排除法定节假日。"],
            ["同日重复", "同一人员同一天只能有一条工时，不同项目也不允许重复。"],
            ["人员成本", "已审批工时 × 审批时人员单天成本，成本后补时可回算 0 成本快照。"],
            ["已发生人天", "某项目已审批通过工时天数合计。"],
            ["预算剩余", "预算金额或人天减去已发生金额或人天，用于预警和项目详情展示。"],
            ["应收未收", "已到期收入合同阶段计划金额减已收金额后的余额。"],
            ["应付未付", "已到期支出合同阶段计划金额减已付金额后的余额。"],
            ["时区", "系统统一使用 Asia/Shanghai；Docker 服务设置 TZ=Asia/Shanghai，后端时间生成和前端展示保持一致。"],
        ],
        widths=[1.6, 5.0],
    )

    heading(doc, "10.2 操作安全建议", 2)
    bullet(doc, "涉及审批、付款、收款和费用的操作，请先进入详情核对项目、合同、阶段、金额、账户和附件。")
    bullet(doc, "导入客户、供应商和合同时，先下载系统模板，保留中文下拉值，避免自行改列名。")
    bullet(doc, "权限调整建议先小范围验证，确认菜单、按钮和数据范围符合预期后再扩大。")
    bullet(doc, "生产环境的飞书 App ID、App Secret、SSH 密码、数据库备份路径等敏感信息不得写入普通用户手册。")


def build_markdown():
    content = f"""# {TITLE}用户使用手册

文档版本：{VERSION}
更新时间：{UPDATE_DATE}

本手册已生成正式 Word 版：`{OUT_DOCX.name}`。Word 版包含完整图文步骤和 29 张页面截图；本 Markdown 文件用于快速查阅和后续维护。

## 阅读说明

- 本版基于 {VERSION} 重新生成。
- 截图位于 `docs/assets/manual-v3.0.56/`，采用常规浏览器视窗比例，避免 4K 缩图问题。
- 系统名称统一为“综合项目管理系统”，英文标识为 “SINYE PROJECT MANAGE SYSTEM”。
- 不记录飞书 App Secret、服务器 SSH 密码等敏感信息。

## 1. 角色速查

| 角色 | 重点功能 | 常见任务 |
|---|---|---|
| 普通员工 | 工作台、我的工时、项目费用 | 填报工时、提交费用或报销、查看审批状态 |
| 工时审批人 | 工时审批、工时详情 | 处理待审批工时，查看审批时间线 |
| 项目负责人 | 工作台项目区、项目详情、项目看板 | 查看负责项目的预算、成本、合同、收付款和预警 |
| 财务/出纳 | 合同收款、合同付款、项目费用、通知中心 | 审核收付款、确认支付、处理失败通知 |
| 系统管理员 | 系统管理 | 维护组织同步、人员成本、角色权限、审批岗位、流程和基础配置 |

## 2. 登录与工作台

![工作台](assets/manual-v3.0.56/01-workbench.png)

1. 在飞书内打开系统应用，或访问管理员提供的系统地址。
2. 点击“使用飞书登录”完成授权。
3. 首次接入时，第一个成功登录的飞书用户会初始化为系统管理员。
4. 登录后，系统按角色、菜单权限、按钮权限和数据范围显示功能。
5. 工作台上半区是项目负责人视角，下半区是员工工时视角。
6. 欢迎语按 Asia/Shanghai 当前时间动态显示，日期行只显示日期和星期。

## 3. 工时管理

### 3.1 我的工时

![我的工时](assets/manual-v3.0.56/02-worklogs.png)

1. 进入“工时管理 > 我的工时”。
2. 按状态筛选草稿、审批中、已驳回或已通过记录。
3. 点击“查看详情”进入单条工时详情；草稿或驳回记录可编辑后重新提交。

### 3.2 新增工时

![新增工时](assets/manual-v3.0.56/03-worklog-form.png)

1. 点击“＋ 新增工时”。
2. 先选择项目类型，再选择该类型下允许填报的项目。
3. 填写开始日期和结束日期，系统自动按周一至周五工作日计算整数天数。
4. 填写项目阶段、工时类型、工作状态和工作描述。
5. 保存草稿或提交审批。

工时规则：

- 同一人员同一天不能填报多次，不同项目也不允许。
- 日期区间不能跨自然月。
- 开始日期不能晚于结束日期。
- 工时单位只有“天”，没有小时。

### 3.3 工时详情与审批

![工时详情](assets/manual-v3.0.56/05-worklog-detail.png)

![工时审批](assets/manual-v3.0.56/04-worklog-approvals.png)

1. 工时详情展示单号、项目、日期、天数、描述、审批人和通知状态。
2. 审批人进入“工时管理 > 工时审批”，在“待我审批”处理工时。
3. 飞书工时审批卡片会直达对应工时审批详情页。
4. 审批通过后，系统更新项目已发生人天，并按审批时人员成本生成项目人员成本。

### 3.4 工时看板与台账

![工时看板](assets/manual-v3.0.56/06-worklog-board.png)

![工时台账](assets/manual-v3.0.56/07-worklog-ledger.png)

- 工时看板只统计工时相关内容，包括部门、项目、人员、趋势和审批状态。
- 部门筛选支持组织层级，选择上级部门时包含下级部门。
- 工时台账是后台维护功能，可按日期、部门、人员和项目筛选，并按权限修改或删除。

## 4. 项目管理

![项目管理](assets/manual-v3.0.56/08-projects.png)

![项目表单](assets/manual-v3.0.56/09-project-form.png)

![项目详情](assets/manual-v3.0.56/10-project-detail.png)

1. 项目列表只展示项目主信息：名称、编号、状态、类型、负责人、所属部门、客户和工时填报开关。
2. 可按项目类型、部门、状态、负责人、客户和标题模糊查询，列表支持分页。
3. 新增或编辑项目在弹窗完成，人员与部门来自飞书组织架构。
4. 项目详情展示概况、基本信息、人员与工时、预算、合同与资金、费用明细和操作记录。
5. 项目预算包含人天预算、成本预算和费用预算；已发生人天由已审批工时自动汇总。

成本口径：

- 实际总成本 = 已确认人员成本 + 已确认合同付款成本 + 已审批项目费用成本。
- 预算剩余 = 对应预算减已发生金额或人天。
- 人员成本按审批通过工时与审批时人员单天成本生成。

## 5. 合同、阶段与收付款

![项目合同](assets/manual-v3.0.56/11-contracts.png)

![合同表单](assets/manual-v3.0.56/12-contract-form.png)

![合同详情](assets/manual-v3.0.56/13-contract-detail.png)

1. 合同分为收入合同和支出合同；实施项目既可能有收入合同，也可能有支出合同。
2. 新建合同可保存草稿，也可在新建时直接上传附件。
3. 点击合同名称进入合同详情，查看主信息、结算率、阶段和款单笔数。
4. 点击“查看阶段”只查看阶段列表；点击“新增阶段”才打开新增阶段弹窗。
5. 合同导入时项目字段填写项目名称，合同类型填写中文，客户/供应商填写名称，系统匹配后自动保存编码。

### 5.1 合同收款单

![合同收款](assets/manual-v3.0.56/14-receipts.png)

![新增合同收款单](assets/manual-v3.0.56/15-receipt-form.png)

1. 选择项目、收入合同和合同阶段。
2. 选择收入合同后，系统自动带出客户付款方名称、付款银行和付款账号。
3. 填写金额、日期、备注并上传附件。
4. 审批通过后通知申请人，增加已收并核减已到期应收。

### 5.2 合同付款单

![合同付款](assets/manual-v3.0.56/16-payments.png)

![新增合同付款单](assets/manual-v3.0.56/17-payment-form.png)

1. 选择项目、支出合同和合同阶段。
2. 选择支出合同后，系统自动带出供应商收款方名称、收款银行和收款账号。
3. 账号、收款方姓名和银行为必填。
4. 合同付款、费用付款和报销款使用同一个付款流程。

## 6. 项目费用

![项目费用](assets/manual-v3.0.56/18-project-expenses.png)

![新增费用](assets/manual-v3.0.56/19-expense-form.png)

1. 费用类型包括差旅费用、无合同发票和其他费用。
2. 新建费用时选择项目、填写金额、日期、事由和说明，并上传附件。
3. 提交后按动态审批流程流转。
4. 审批通过后计入项目费用成本；出纳支付后计入已付。

## 审批流程摘要

- 差旅费用：经办人提交 → 项目负责人审批 → 会计审批 → 上级领导审批 → 财务总监审批 → 总经理审批 → 出纳支付。
- 费用付款 / 无合同发票 / 其他费用 / 合同付款：经办人提交 → 项目负责人审批 → 会计审批 → 经营审批 → 上级领导审批 → 经营副总审批 → 财务总监审批 → 总经理审批 → 出纳支付。

## 7. 项目看板与提醒

![项目看板](assets/manual-v3.0.56/21-project-board.png)

项目看板支持“我的项目”和“全部项目”范围。选择全部项目时，项目选择框会动态获取当前可访问项目列表。

| 指标 | 说明 |
|---|---|
| 人员成本 | 已审批工时 × 审批时人员单天成本 |
| 合同成本 | 支出合同付款确认后的成本 |
| 项目费用 | 差旅、无合同发票和其他费用审批通过后的成本 |
| 应收未收 | 只统计已到期收入阶段的未收金额 |
| 应付未付 | 只统计已到期支出阶段的未付金额 |

系统提醒包括工时待审批、合同收款单待审批、合同付款单待审批、项目费用待审批、合同阶段临期/到期/逾期、成本预算接近或超支、人天预算接近或超支。飞书通知只负责提醒和打开系统页面，业务状态以本系统为准。

## 8. 客户、供应商、人员成本

![供应商管理](assets/manual-v3.0.56/22-suppliers.png)

![客户管理](assets/manual-v3.0.56/23-customers.png)

![人员成本](assets/manual-v3.0.56/24-personnel-costs.png)

- 供应商和客户编码由系统自动生成，导入模板不需要填写编码。
- 供应商维护默认收款账户，合同付款单选择支出合同后可自动带出。
- 客户维护默认付款账户，合同收款单选择收入合同后可自动带出。
- 人员成本按部门和人员维护；部门筛选支持层级包含。

## 关键口径摘要

- 工时单位为天，只统计周一至周五工作日。
- 同一人员同一天不能重复填报工时。
- 人员成本按已审批工时和审批时人员单天成本生成。
- 应收未收/应付未付只统计已到期阶段余额。
- 系统统一使用 Asia/Shanghai 时区。

## 9. 系统管理

![集成与通知](assets/manual-v3.0.56/25-integrations.png)

![通知中心](assets/manual-v3.0.56/26-notification-center.png)

![管理员管理](assets/manual-v3.0.56/27-admin-permissions.png)

![审批岗位](assets/manual-v3.0.56/28-approval-positions.png)

![审批流程](assets/manual-v3.0.56/29-approval-workflows.png)

![基础配置](assets/manual-v3.0.56/30-configuration.png)

1. 集成与通知：查看飞书连接、同步组织架构、统一维护通知策略。
2. 通知中心：查看发送记录、失败原因和重试入口。
3. 管理员管理：配置角色、菜单、按钮和数据访问范围。
4. 审批岗位：维护会计、经营审批、财务总监、总经理、出纳等岗位成员。
5. 审批流程：查看各单据类型的流程节点和版本。
6. 基础配置：维护项目类型、项目状态、工时类型、项目阶段等动态下拉项。

## 10. 常见问题

| 问题 | 处理建议 |
|---|---|
| 看不到项目 | 检查是否为项目负责人、数据范围是否授权、部门是否同步、筛选条件是否限制 |
| 看不到菜单 | 检查角色菜单权限 |
| 工时提交失败 | 检查是否同日重复、跨月、项目不可填报或天数不是整数工作日 |
| 飞书没有通知 | 在通知中心查看通知记录、发送状态、失败原因和接收人 open_id |
| 导出没有下载弹窗 | 查看飞书内置浏览器下载栏或下载记录 |
| 附件预览失败 | Office 文件转换失败时可下载本地查看 |
"""
    OUT_MD.write_text(content, encoding="utf-8")
    CANONICAL_MD.write_text(content, encoding="utf-8")


def build_docx():
    doc = Document()
    configure_doc(doc)
    cover(doc)
    add_revision_summary(doc)
    add_contents(doc)
    doc.add_page_break()
    add_login_and_workbench(doc)
    doc.add_page_break()
    add_worklog(doc)
    doc.add_page_break()
    add_project(doc)
    doc.add_page_break()
    add_contract_and_cash(doc)
    doc.add_page_break()
    add_expenses(doc)
    doc.add_page_break()
    add_boards(doc)
    doc.add_page_break()
    add_counterparties_and_cost(doc)
    doc.add_page_break()
    add_system_admin(doc)
    doc.add_page_break()
    add_faq_and_data(doc)
    paragraph(doc, "", after=0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("— 文档结束 —")
    set_font(r, size=10, color=MUTED, italic=True)
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)


def main():
    build_markdown()
    build_docx()
    print("done")


if __name__ == "__main__":
    main()

